#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
fill_proposal.py — Điền thông tin nhóm + phân công công việc vào Proposal đồ án.

Lấy file gốc "Website ban do noi that.docx" (template của nhóm), điền:
  1. Vai trò 3 thành viên ở mục "Thông tin chung".
  2. Toàn bộ placeholder "…………" trong các bảng phân công 10 tuần.
Rồi lưu thành docs/Proposal-WebNoiThat.docx trong repo.
Kèm bổ sung thông tin nhóm vào docs/BaoCao-WebNoiThat.docx.

Chạy: python scripts/fill_proposal.py
Yêu cầu: python-docx
"""
import shutil
from pathlib import Path

from docx import Document

BASE = Path(r"D:\WORK\NAM4\3_CDTH")
MASTER = BASE / "Website ban do noi that.docx"
REPO = BASE / "Website_NoiThat"
OUT_PROPOSAL = REPO / "docs" / "Proposal-WebNoiThat.docx"
OUT_REPORT = REPO / "docs" / "BaoCao-WebNoiThat.docx"

# ===== Thông tin nhóm (nguồn chuẩn duy nhất) =====
MEMBERS = [
    ("Trần Đức Lương", "1502368", "Nhóm trưởng · Backend"),
    ("Nguyễn Như Kiên", "0015768", "Backend"),
    ("Phan Thế Kiệt", "0015968", "Frontend"),
]
GVHD = "Phạm Hữu Tùng"

# Bản đồ: dòng placeholder trong cột "Phân công công việc" -> người phụ trách.
# Key = phần đầu dòng trước dấu ':', khớp bằng startswith.
ASSIGNMENTS = {
    # Tuần 1–2: Analysis
    "Viết BRS": MEMBERS[0][0],
    "Viết User Story": MEMBERS[1][0],
    "Vẽ Use Case Diagram": MEMBERS[1][0],
    "Phân tích flow đặt hàng": MEMBERS[0][0],
    "Research hệ thống tương tự (Shopee, Tiki nội thất)": MEMBERS[2][0],
    "Viết SRS + vẽ flowchart": MEMBERS[1][0],
    # Tuần 3–4: Design
    "Thiết kế UI (Figma)": MEMBERS[2][0],
    "Kiểm soát thiết kế tổng thể": MEMBERS[0][0],
    "Thiết kế Database (ERD)": MEMBERS[0][0],
    "Class + Sequence Diagram": MEMBERS[1][0],
    "Wireframe + hỗ trợ UI": MEMBERS[2][0],
    "Activity Diagram": MEMBERS[1][0],
    # Tuần 5–6: Setup
    "Quản lý GitHub repo, merge code": MEMBERS[0][0],
    "Setup Database, Migration, cấu trúc backend": MEMBERS[1][0],
    "Setup Frontend (Blade/Bootstrap)": MEMBERS[2][0],
    "Test môi trường + fix lỗi": MEMBERS[0][0],
    "Seed dữ liệu mẫu": MEMBERS[1][0],
    # Tuần 7–8: Coding
    "Giỏ hàng, đặt hàng, logic nghiệp vụ": MEMBERS[0][0],
    "Auth + User": MEMBERS[1][0],
    # Dòng gốc ghi nhầm "Room, Booking" (sót từ template mẫu khách sạn) — đã sửa qua TEXT_FIXES
    "Backend API (Products, Orders, Cart)": MEMBERS[1][0],
    "Giao diện khách hàng + admin dashboard": MEMBERS[2][0],
    "Quản lý đơn + user (admin)": MEMBERS[0][0],
    # Tuần 9–10: Hoàn thiện
    "Tổng hợp báo cáo + thuyết trình": MEMBERS[0][0],
    "Fix backend bug + testing": MEMBERS[1][0],
    "Fix UI/UX": MEMBERS[2][0],
    "Deploy + demo video": MEMBERS[0][0],
}

# Sửa chữ sót lại từ template mẫu "Đặt phòng khách sạn"
TEXT_FIXES = {
    "Backend API (Room, Booking, Products, Orders)": "Backend API (Products, Orders, Cart)",
}


def fill_members(doc: Document) -> None:
    """Điền vai trò vào danh sách thành viên ở mục Thông tin chung."""
    for para in doc.paragraphs:
        t = para.text.strip()
        for name, mssv, role in MEMBERS:
            if name in t and str(mssv) in t.replace(" ", ""):
                para.text = f"{name} – {mssv} ({role})"
                break
        else:
            if "Giảng viên hướng dẫn" in t and not GVHD in t:
                para.text = f"Giảng viên hướng dẫn: {GVHD}"


def fill_assignment_tables(doc: Document) -> int:
    """Điền tên người phụ trách vào mọi dòng '…………' trong bảng phân công. Trả về số dòng đã điền."""
    filled = 0
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 3:
                continue
            cell = row.cells[2]
            for para in cell.paragraphs:
                raw = para.text
                # Mỗi ô phân công là MỘT paragraph chứa nhiều dòng ngắt mềm (w:br).
                lines = raw.split("\n")
                if "…" not in raw or ":" not in raw:
                    continue
                new_lines: list[str] = []
                for line in lines:
                    stripped = line.strip()
                    if "…" in stripped and ":" in stripped:
                        key = stripped.split(":")[0].strip()
                        key = TEXT_FIXES.get(key, key)
                        owner = ASSIGNMENTS.get(key)
                        if owner:
                            line = f"{key}: {owner}"
                            filled += 1
                        else:
                            print(f"  [CẢNH BÁO] Chưa khớp key: {key!r}")
                    new_lines.append(line)
                # Dựng lại paragraph: xóa hết run cũ, ghi dòng mới với add_break giữa các dòng
                for run in list(para.runs):
                    run.text = ""
                first = para.runs[0] if para.runs else para.add_run("")
                for i, nl in enumerate(new_lines):
                    if i > 0:
                        first.add_break()
                    first.add_text(nl)
    return filled


def patch_report(doc: Document) -> None:
    """Bổ sung khối thông tin nhóm vào đầu Báo cáo (sau tiêu đề).

    insert_paragraph_before luôn chèn NGAY TRƯỚC anchor nên phải duyệt XUÔI
    để giữ đúng thứ tự các dòng.
    """
    # Đã có khối thông tin nhóm ở đầu thì bỏ qua (chạy lại an toàn)
    joined = "\n".join(p.text for p in doc.paragraphs[:8])
    if "Nhóm thực hiện" in joined:
        return
    anchor = doc.paragraphs[1]  # đoạn ngay dưới Heading chính
    lines = ["Nhóm thực hiện (3 thành viên):"]
    lines += [f"• {n} – {m} ({r})" for n, m, r in MEMBERS]
    lines.append(f"Giảng viên hướng dẫn: {GVHD}")
    for line in lines:
        anchor.insert_paragraph_before(line)


def main() -> None:
    OUT_PROPOSAL.parent.mkdir(parents=True, exist_ok=True)

    # 1) Proposal: copy từ file gốc rồi điền
    shutil.copyfile(MASTER, OUT_PROPOSAL)
    prop = Document(str(OUT_PROPOSAL))
    fill_members(prop)
    n = fill_assignment_tables(prop)
    prop.save(str(OUT_PROPOSAL))
    print(f"Đã điền {n} dòng phân công -> {OUT_PROPOSAL}")

    # 2) Báo cáo: bổ sung thông tin nhóm nếu chưa có
    rep = Document(str(OUT_REPORT))
    patch_report(rep)
    rep.save(str(OUT_REPORT))
    print(f"Đã cập nhật thông tin nhóm -> {OUT_REPORT}")


if __name__ == "__main__":
    main()
